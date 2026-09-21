from typing import Dict, Any
from jinja2 import Template as JinjaTemplate
import weasyprint
import io
import docx

class DocumentGenerationService:
    @staticmethod
    def generate_html(template_html: str, variables: Dict[str, Any]) -> str:
        """Interpolate variables into HTML using Jinja2."""
        jinja_template = JinjaTemplate(template_html)
        return jinja_template.render(**variables)

    @staticmethod
    def generate_pdf(template_html: str, variables: Dict[str, Any]) -> bytes:
        """Generate PDF from HTML and variables using weasyprint."""
        html_content = DocumentGenerationService.generate_html(template_html, variables)
        pdf_bytes = weasyprint.HTML(string=html_content).write_pdf()
        return pdf_bytes

    @staticmethod
    def generate_docx(template_html: str, variables: Dict[str, Any]) -> bytes:
        """Generate DOCX. In a real system, you might map HTML to DOCX, 
        or use a pre-existing DOCX template. Here we just create a simple DOCX 
        with the interpolated HTML as text to provide scaffolding."""
        html_content = DocumentGenerationService.generate_html(template_html, variables)
        
        doc = docx.Document()
        doc.add_paragraph("Document generated from template:")
        doc.add_paragraph(html_content)
        
        file_stream = io.BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)
        return file_stream.read()
